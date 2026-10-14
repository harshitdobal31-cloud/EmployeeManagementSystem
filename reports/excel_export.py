from pathlib import Path
import csv

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    Alignment,
    PatternFill,
    Border,
    Side,
)
from openpyxl.utils import get_column_letter


# ============================================================
# PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent

REPORTS_DIR = BASE_DIR / "reports"

EXPORT_DIR = REPORTS_DIR / "exports"

EXPORT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


# ============================================================
# COLORS
# ============================================================

NAVY = "0F172A"
BLUE = "2563EB"
LIGHT_BLUE = "DBEAFE"
WHITE = "FFFFFF"
BORDER_COLOR = "D9E2F0"


# ============================================================
# EXCEL EXPORT
# ============================================================

def export_to_excel(
    filename,
    title,
    headers,
    rows,
):
    """
    Create a professional Excel report.

    Returns:
        Path object of generated Excel file.
    """

    try:

        if not filename.lower().endswith(".xlsx"):
            filename += ".xlsx"

        file_path = EXPORT_DIR / filename

        workbook = Workbook()

        worksheet = workbook.active

        worksheet.title = "Report"

        # ----------------------------------------------------
        # TITLE
        # ----------------------------------------------------

        column_count = max(len(headers), 1)

        worksheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=column_count,
        )

        title_cell = worksheet.cell(
            row=1,
            column=1,
        )

        title_cell.value = title

        title_cell.font = Font(
            name="Calibri",
            size=18,
            bold=True,
            color=WHITE,
        )

        title_cell.fill = PatternFill(
            fill_type="solid",
            fgColor=NAVY,
        )

        title_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        worksheet.row_dimensions[1].height = 34

        # ----------------------------------------------------
        # GENERATED INFO
        # ----------------------------------------------------

        worksheet.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=column_count,
        )

        info_cell = worksheet.cell(
            row=2,
            column=1,
        )

        from datetime import datetime

        info_cell.value = (
            "Employee Management System | "
            f"Generated: {datetime.now().strftime('%d-%m-%Y %I:%M %p')}"
        )

        info_cell.font = Font(
            size=10,
            color="64748B",
        )

        info_cell.alignment = Alignment(
            horizontal="center",
        )

        # ----------------------------------------------------
        # HEADER
        # ----------------------------------------------------

        header_row = 4

        header_fill = PatternFill(
            fill_type="solid",
            fgColor=BLUE,
        )

        header_font = Font(
            name="Calibri",
            size=11,
            bold=True,
            color=WHITE,
        )

        thin_side = Side(
            style="thin",
            color=BORDER_COLOR,
        )

        border = Border(
            left=thin_side,
            right=thin_side,
            top=thin_side,
            bottom=thin_side,
        )

        for column_index, header in enumerate(
            headers,
            start=1,
        ):

            cell = worksheet.cell(
                row=header_row,
                column=column_index,
            )

            cell.value = header

            cell.fill = header_fill

            cell.font = header_font

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            cell.border = border

        worksheet.row_dimensions[
            header_row
        ].height = 28

        # ----------------------------------------------------
        # DATA
        # ----------------------------------------------------

        for row_index, row_data in enumerate(
            rows,
            start=header_row + 1,
        ):

            for column_index, value in enumerate(
                row_data,
                start=1,
            ):

                cell = worksheet.cell(
                    row=row_index,
                    column=column_index,
                )

                cell.value = value

                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=True,
                )

                cell.border = border

                # Currency
                if isinstance(value, float):
                    cell.number_format = '#,##0.00'

        # ----------------------------------------------------
        # COLUMN WIDTH
        # ----------------------------------------------------

        for column_index in range(
            1,
            column_count + 1,
        ):

            column_letter = get_column_letter(
                column_index
            )

            maximum_length = 0

            for cell in worksheet[
                column_letter
            ]:

                if cell.value is not None:

                    length = len(
                        str(cell.value)
                    )

                    maximum_length = max(
                        maximum_length,
                        length,
                    )

            width = min(
                max(
                    maximum_length + 3,
                    14,
                ),
                42,
            )

            worksheet.column_dimensions[
                column_letter
            ].width = width

        # ----------------------------------------------------
        # FREEZE
        # ----------------------------------------------------

        worksheet.freeze_panes = "A5"

        # ----------------------------------------------------
        # FILTER
        # ----------------------------------------------------

        if rows and headers:

            last_row = (
                header_row
                + len(rows)
            )

            last_column = get_column_letter(
                len(headers)
            )

            worksheet.auto_filter.ref = (
                f"A{header_row}:"
                f"{last_column}{last_row}"
            )

        # ----------------------------------------------------
        # PRINT SETTINGS
        # ----------------------------------------------------

        worksheet.sheet_view.showGridLines = False

        worksheet.page_setup.orientation = "landscape"

        worksheet.page_setup.fitToWidth = 1

        worksheet.page_setup.fitToHeight = 0

        worksheet.sheet_properties.pageSetUpPr.fitToPage = True

        # ----------------------------------------------------
        # SAVE
        # ----------------------------------------------------

        workbook.save(file_path)

        return file_path

    except Exception as error:

        raise Exception(
            f"Unable to export Excel report: {error}"
        )


# ============================================================
# CSV EXPORT
# ============================================================

def export_to_csv(
    filename,
    headers,
    rows,
):
    """
    Export report to CSV.
    """

    try:

        if not filename.lower().endswith(".csv"):
            filename += ".csv"

        file_path = EXPORT_DIR / filename

        with open(
            file_path,
            "w",
            newline="",
            encoding="utf-8-sig",
        ) as csv_file:

            writer = csv.writer(
                csv_file
            )

            writer.writerow(headers)

            writer.writerows(rows)

        return file_path

    except Exception as error:

        raise Exception(
            f"Unable to export CSV report: {error}"
        )


# ============================================================
# EXPORT DIRECTORY
# ============================================================

def get_export_directory():

    EXPORT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    return EXPORT_DIR