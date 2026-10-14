from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent.parent

EXPORT_DIR = (
    BASE_DIR
    / "reports"
    / "payroll_exports"
)

EXPORT_DIR.mkdir(
    parents=True,
    exist_ok=True
)


def export_payroll_to_excel(
    payroll_records,
    file_path=None
):
    """
    Export payroll records to an Excel file.

    payroll_records format:

        [
            (payroll, employee),
            (payroll, employee),
            ...
        ]
    """

    # --------------------------------------------------------
    # Default file path
    # --------------------------------------------------------

    if file_path is None:

        file_path = (
            EXPORT_DIR
            / "Payroll_Report.xlsx"
        )

    else:

        file_path = Path(
            file_path
        )

    # --------------------------------------------------------
    # Workbook
    # --------------------------------------------------------

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Payroll Report"

    # --------------------------------------------------------
    # Colors
    # --------------------------------------------------------

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="0F2747"
    )

    subheader_fill = PatternFill(
        fill_type="solid",
        fgColor="E8F1FF"
    )

    total_fill = PatternFill(
        fill_type="solid",
        fgColor="DBEAFE"
    )

    white_font = Font(
        color="FFFFFF",
        bold=True,
        size=11
    )

    normal_font = Font(
        color="1F2937",
        size=10
    )

    title_font = Font(
        bold=True,
        size=18,
        color="0F2747"
    )

    bold_font = Font(
        bold=True,
        size=10,
        color="0F2747"
    )

    thin_side = Side(
        style="thin",
        color="D1D5DB"
    )

    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side
    )

    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------

    worksheet.merge_cells(
        "A1:K1"
    )

    worksheet["A1"] = (
        "EMPLOYEE MANAGEMENT SYSTEM"
    )

    worksheet["A1"].font = title_font

    worksheet["A1"].alignment = (
        Alignment(
            horizontal="center",
            vertical="center"
        )
    )

    worksheet.row_dimensions[1].height = 30

    # --------------------------------------------------------
    # Subtitle
    # --------------------------------------------------------

    worksheet.merge_cells(
        "A2:K2"
    )

    worksheet["A2"] = (
        "Payroll Report"
    )

    worksheet["A2"].font = Font(
        size=12,
        color="64748B",
        italic=True
    )

    worksheet["A2"].alignment = (
        Alignment(
            horizontal="center"
        )
    )

    # --------------------------------------------------------
    # Headers
    # --------------------------------------------------------

    headers = [
        "Payroll ID",
        "Employee ID",
        "Employee Name",
        "Department",
        "Designation",
        "Payroll Month",
        "Basic Salary",
        "Allowances",
        "Deductions",
        "Net Salary",
        "Status",
    ]

    header_row = 4

    for column, header in enumerate(
        headers,
        start=1
    ):

        cell = worksheet.cell(
            row=header_row,
            column=column
        )

        cell.value = header

        cell.fill = header_fill

        cell.font = white_font

        cell.alignment = (
            Alignment(
                horizontal="center",
                vertical="center"
            )
        )

        cell.border = border

    worksheet.row_dimensions[
        header_row
    ].height = 25

    # --------------------------------------------------------
    # Data
    # --------------------------------------------------------

    starting_row = 5

    total_basic = 0
    total_allowances = 0
    total_deductions = 0
    total_net = 0

    for index, record in enumerate(
        payroll_records,
        start=starting_row
    ):

        payroll = record[0]

        employee = record[1]

        basic_salary = float(
            payroll.basic_salary or 0
        )

        allowances = float(
            payroll.allowances or 0
        )

        deductions = float(
            payroll.deductions or 0
        )

        net_salary = float(
            payroll.net_salary or 0
        )

        total_basic += basic_salary
        total_allowances += allowances
        total_deductions += deductions
        total_net += net_salary

        row_data = [
            payroll.id,
            employee.employee_id,
            employee.full_name,
            employee.department or "-",
            employee.designation or "-",
            payroll.payroll_month,
            basic_salary,
            allowances,
            deductions,
            net_salary,
            payroll.status,
        ]

        for column, value in enumerate(
            row_data,
            start=1
        ):

            cell = worksheet.cell(
                row=index,
                column=column
            )

            cell.value = value

            cell.font = normal_font

            cell.border = border

            cell.alignment = (
                Alignment(
                    vertical="center"
                )
            )

        # Currency columns

        for column in [
            7,
            8,
            9,
            10
        ]:

            worksheet.cell(
                row=index,
                column=column
            ).number_format = (
                '₹ #,##0.00'
            )

            worksheet.cell(
                row=index,
                column=column
            ).alignment = (
                Alignment(
                    horizontal="right"
                )
            )

        # Center selected columns

        for column in [
            1,
            2,
            6,
            11
        ]:

            worksheet.cell(
                row=index,
                column=column
            ).alignment = (
                Alignment(
                    horizontal="center",
                    vertical="center"
                )
            )

    # --------------------------------------------------------
    # Total Row
    # --------------------------------------------------------

    total_row = (
        starting_row
        + len(payroll_records)
    )

    worksheet.cell(
        row=total_row,
        column=1
    ).value = "TOTAL"

    worksheet.cell(
        row=total_row,
        column=1
    ).font = bold_font

    worksheet.cell(
        row=total_row,
        column=1
    ).fill = total_fill

    worksheet.merge_cells(
        start_row=total_row,
        start_column=1,
        end_row=total_row,
        end_column=6
    )

    worksheet.cell(
        row=total_row,
        column=7
    ).value = total_basic

    worksheet.cell(
        row=total_row,
        column=8
    ).value = total_allowances

    worksheet.cell(
        row=total_row,
        column=9
    ).value = total_deductions

    worksheet.cell(
        row=total_row,
        column=10
    ).value = total_net

    for column in range(
        1,
        12
    ):

        cell = worksheet.cell(
            row=total_row,
            column=column
        )

        cell.fill = total_fill

        cell.font = bold_font

        cell.border = border

    for column in [
        7,
        8,
        9,
        10
    ]:

        worksheet.cell(
            row=total_row,
            column=column
        ).number_format = (
            '₹ #,##0.00'
        )

        worksheet.cell(
            row=total_row,
            column=column
        ).alignment = (
            Alignment(
                horizontal="right"
            )
        )

    # --------------------------------------------------------
    # Freeze panes
    # --------------------------------------------------------

    worksheet.freeze_panes = "A5"

    # --------------------------------------------------------
    # Auto filter
    # --------------------------------------------------------

    worksheet.auto_filter.ref = (
        f"A4:K{total_row}"
    )

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    widths = {
        "A": 13,
        "B": 16,
        "C": 25,
        "D": 20,
        "E": 24,
        "F": 18,
        "G": 17,
        "H": 17,
        "I": 17,
        "J": 17,
        "K": 15,
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    workbook.save(
        file_path
    )

    return file_path