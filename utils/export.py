from pathlib import Path
import csv

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# REPORT EXPORT DIRECTORY
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent
REPORTS_DIR = BASE_DIR / "reports"
EXPORT_DIR = REPORTS_DIR / "exports"

EXPORT_DIR.mkdir(parents=True, exist_ok=True)


# ============================================================
# EXCEL EXPORT
# ============================================================

def export_to_excel(
    filename,
    title,
    headers,
    rows,
):
    """
    Export report data to a professional Excel file.

    Parameters:
        filename : Output filename
        title    : Report title
        headers  : List of column headings
        rows     : List of row data

    Returns:
        Path of generated Excel file
    """

    try:
        if not filename.lower().endswith(".xlsx"):
            filename += ".xlsx"

        file_path = EXPORT_DIR / filename

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Report"

        # ----------------------------------------------------
        # TITLE
        # ----------------------------------------------------

        worksheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=len(headers),
        )

        title_cell = worksheet.cell(row=1, column=1)
        title_cell.value = title
        title_cell.font = Font(
            bold=True,
            size=16,
        )
        title_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        worksheet.row_dimensions[1].height = 30

        # ----------------------------------------------------
        # HEADER
        # ----------------------------------------------------

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        header_font = Font(
            bold=True,
            color="FFFFFF",
        )

        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        header_row = 3

        for column_index, header in enumerate(headers, start=1):

            cell = worksheet.cell(
                row=header_row,
                column=column_index,
            )

            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
            cell.border = thin_border

        # ----------------------------------------------------
        # DATA
        # ----------------------------------------------------

        data_start_row = header_row + 1

        for row_index, row_data in enumerate(
            rows,
            start=data_start_row,
        ):

            for column_index, value in enumerate(
                row_data,
                start=1,
            ):

                cell = worksheet.cell(
                    row=row_index,
                    column=column_index,
                )

                cell.value = value

                cell.alignment = Alignment(
                    vertical="center",
                )

                cell.border = thin_border

        # ----------------------------------------------------
        # AUTO COLUMN WIDTH
        # ----------------------------------------------------

        for column_index in range(1, len(headers) + 1):

            column_letter = get_column_letter(
                column_index
            )

            maximum_length = 0

            for cell in worksheet[column_letter]:

                if cell.value is not None:

                    cell_length = len(
                        str(cell.value)
                    )

                    if cell_length > maximum_length:
                        maximum_length = cell_length

            adjusted_width = min(
                maximum_length + 3,
                40,
            )

            worksheet.column_dimensions[
                column_letter
            ].width = max(
                adjusted_width,
                12,
            )

        # ----------------------------------------------------
        # FREEZE HEADER
        # ----------------------------------------------------

        worksheet.freeze_panes = "A4"

        # ----------------------------------------------------
        # FILTER
        # ----------------------------------------------------

        if rows:
            worksheet.auto_filter.ref = (
                f"A3:{get_column_letter(len(headers))}"
                f"{len(rows) + 3}"
            )

        # ----------------------------------------------------
        # SAVE
        # ----------------------------------------------------

        workbook.save(file_path)

        return file_path

    except Exception as error:

        raise Exception(
            f"Unable to export Excel report: {error}"
        )


# ============================================================
# CSV EXPORT
# ============================================================

def export_to_csv(
    filename,
    headers,
    rows,
):
    """
    Export report data to CSV.

    Parameters:
        filename : Output filename
        headers  : List of column headings
        rows     : List of row data

    Returns:
        Path of generated CSV file
    """

    try:

        if not filename.lower().endswith(".csv"):
            filename += ".csv"

        file_path = EXPORT_DIR / filename

        with open(
            file_path,
            mode="w",
            newline="",
            encoding="utf-8-sig",
        ) as csv_file:

            writer = csv.writer(csv_file)

            # Header
            writer.writerow(headers)

            # Data
            writer.writerows(rows)

        return file_path

    except Exception as error:

        raise Exception(
            f"Unable to export CSV report: {error}"
        )


# ============================================================
# EXPORT HELPER
# ============================================================

def get_export_directory():
    """
    Return the report export directory.
    """

    EXPORT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    return EXPORT_DIR